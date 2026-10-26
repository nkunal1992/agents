"""extract_files: real bytes → chunks per format; warnings, never crashes."""
from __future__ import annotations

import io

import pytest

from confluence_knowledge.extract_files import (
    extract_file_chunks,
    is_secret_name,
)

pytest.importorskip("docx")
pytest.importorskip("openpyxl")
pytest.importorskip("pptx")
pytest.importorskip("pypdf")


def _docx_bytes(text: str = "The deploy runbook says restart the gateway first.") -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Deploy Runbook", level=1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Owners"
    ws.append(["service", "owner"])
    ws.append(["gateway", "narayan"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Migration Plan"
    body = slide.placeholders[1]
    body.text = "Phase one: move the database."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_txt_chunks():
    chunks, warns = extract_file_chunks("notes.txt", b"hello world " * 100)
    assert chunks and not warns
    assert "hello world" in chunks[0].text


def test_md_chunks():
    chunks, _ = extract_file_chunks("r.md", b"# Title\n\nsome content")
    assert chunks and "content" in chunks[0].text


def test_csv_chunks():
    chunks, _ = extract_file_chunks("t.csv", b"a,b\n1,2\n3,4\n")
    assert chunks[0].kind == "table" and "row 1: 1 | 2" in chunks[0].text


def test_html_chunks():
    chunks, _ = extract_file_chunks("p.html", b"<h1>H</h1><p>body text</p>")
    assert chunks and "body text" in chunks[0].text


def test_docx_chunks():
    chunks, warns = extract_file_chunks("runbook.docx", _docx_bytes())
    assert chunks and not warns
    assert "gateway" in " ".join(c.text for c in chunks)


def test_docx_heading_path():
    chunks, _ = extract_file_chunks("r.docx", _docx_bytes())
    assert any("Deploy Runbook" in c.heading_path for c in chunks)


def test_xlsx_chunks():
    chunks, _ = extract_file_chunks("o.xlsx", _xlsx_bytes())
    text = " ".join(c.text for c in chunks)
    assert "gateway" in text and "narayan" in text
    assert any("Sheet: Owners" in c.heading_path for c in chunks)


def test_pptx_chunks():
    chunks, _ = extract_file_chunks("p.pptx", _pptx_bytes())
    text = " ".join(c.text for c in chunks)
    assert "Migration Plan" in text and "database" in text


def test_binary_unknown_type_warns():
    """#281: real binaries with unknown extensions are disclosed, not
    treated as text — OLE office, zips and images must never index."""
    png = (b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + b"\x00" * 100
           + b"\x89\x50\x4e\x47" * 200)
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512 + b"garbage"
    zipf = b"PK\x03\x04\x14\x00\x06\x00\x08\x00" + b"\x00\x01" * 300
    for name, blob in (("img.png", png), ("old.doc", ole), ("x.bin", zipf)):
        chunks, warns = extract_file_chunks(name, blob)
        assert not chunks and warns and "binary" in warns[0], \
            f"{name}: warns={warns} chunks={len(chunks)}"


def test_utf16_text_autodetects():
    """#281: UTF-16 files are NUL-dense in bytes but ARE text — the
    alternating-byte pattern must keep them indexable (BOM and BOM-less)."""
    body = "deploy runbook: restart the gateway first\n" * 20
    for name, blob in (("notes.txt", body.encode("utf-16")),
                       ("n.log", body.encode("utf-16-le")),
                       ("n.log", body.encode("utf-16-be"))):
        chunks, warns = extract_file_chunks(name, blob)
        assert chunks and not warns, f"{name}/{len(blob)}: {warns}"
        assert "gateway" in " ".join(c.text for c in chunks)


def test_unknown_text_extension_autodetects():
    """#281: ordinary text files outside the allowlist (logs, configs,
    source) must index — the extension list is not a coverage gate."""
    for name in ("app.log", "conf.yaml", "notes.json", "run.py",
                 "Makefile", "README", "data.tsv~"):
        chunks, warns = extract_file_chunks(
            name, b"deploy runbook: restart the gateway first\n" * 20)
        assert chunks and not warns, f"{name}: {warns}"
        assert "gateway" in " ".join(c.text for c in chunks)


def test_tsv_chunks():
    chunks, _ = extract_file_chunks("t.tsv", b"a\tb\n1\t2\n3\t4\n")
    assert chunks[0].kind == "table" and "row 1: 1 | 2" in chunks[0].text


def test_empty_file_warns():
    chunks, warns = extract_file_chunks("empty.txt", b"")
    assert not chunks and warns and "empty" in warns[0]


def test_large_docx_not_skipped_by_byte_estimate():
    """#281: the old `len(data)//6` guard skipped any file over ~7.2MB —
    a docx is zip-compressed, so embedded media bytes counted as 'words'.
    Office formats are bounded by chunk count, not raw bytes."""
    import io
    import zipfile

    import docx
    d = docx.Document()
    d.add_paragraph("deploy runbook: restart the gateway first")
    buf = io.BytesIO()
    d.save(buf)
    # pad the zip with a large inert media part — 9MB of noise bytes
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as zout:
        for item in zin.namelist():
            zout.writestr(item, zin.read(item))
        zout.writestr("word/media/big.bin", b"\x07" * (9 * 1024 * 1024))
    chunks, warns = extract_file_chunks("runbook.docx", out.getvalue())
    assert chunks and "gateway" in " ".join(c.text for c in chunks), warns


def test_corrupt_office_warns_not_crashes():
    chunks, warns = extract_file_chunks("bad.docx", b"not a zip")
    assert not chunks and warns


def test_secret_names():
    assert is_secret_name(".env") and is_secret_name("id_rsa")
    assert is_secret_name("cert.pem") and is_secret_name("credentials.json")
    assert is_secret_name(".env.production")
    assert not is_secret_name("deploy.docx") and not is_secret_name("notes.txt")


def test_pdf_minimal():
    """A real (hand-built) one-page PDF with extractable text."""
    pdf = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 70>>stream
BT /F1 24 Tf 100 700 Td (deploy steps live here) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
"""
    chunks, warns = extract_file_chunks("doc.pdf", pdf)
    # pypdf extracts the text; if a version can't parse the hand-rolled xref,
    # it must warn — never crash
    assert chunks or warns
    if chunks:
        assert "deploy steps" in " ".join(c.text for c in chunks)
